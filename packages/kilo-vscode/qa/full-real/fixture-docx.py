#!/usr/bin/env python3

import argparse
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


def field(paragraph, code):
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = code
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    placeholder = OxmlElement("w:t")
    placeholder.text = "目录将在渲染时刷新"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run = paragraph.add_run()._r
    for node in (begin, instruction, separate, placeholder, end):
        run.append(node)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--output", required=True)
    parser.add_argument("--run", default="manual")
    parser.add_argument("--no-toc", action="store_true")
    args = parser.parse_args()
    path = Path(args.output).resolve()
    path.parent.mkdir(parents=True, exist_ok=True)

    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)
    for name in ("Normal", "Title", "Subtitle", "Heading 1", "Heading 2", "Heading 3"):
        style = doc.styles[name]
        style.font.name = "STSong"
        style._element.rPr.rFonts.set(qn("w:ascii"), "STSong")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "STSong")
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "STSong")
    doc.styles["Normal"].font.size = Pt(10.5)

    title = doc.add_heading("ChipMate 真实渲染 QA 夹具", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle = doc.add_paragraph(f"Run ID: {args.run}")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if not args.no_toc:
        toc_title = doc.add_paragraph()
        toc_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        toc_run = toc_title.add_run("目录")
        toc_run.bold = True
        toc_run.font.size = Pt(14)
        toc_style = doc.styles.add_style("TOC 1", WD_STYLE_TYPE.PARAGRAPH)
        toc_style.base_style = doc.styles["Normal"]
        field(doc.add_paragraph(style=toc_style), 'TOC \\o "1-3" \\h \\z \\u')
        doc.add_page_break()
    doc.add_heading("一、目的", level=1)
    doc.add_paragraph("此文档只包含固定、无敏感信息的中文内容，用于验证 DOCX create、field refresh、PDF 与 PNG 渲染。")
    doc.add_heading("二、确定性标记", level=1)
    doc.add_paragraph(f"QA-WORD-{args.run}")
    doc.add_heading("三、验收表", level=1)
    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    for cell, text in zip(table.rows[0].cells, ("项目", "预期", "状态")):
        cell.text = text
    for item, expected in (("中文", "保持可读"), ("目录", "字段可刷新"), ("输出", "PDF 与逐页 PNG")):
        cells = table.add_row().cells
        cells[0].text = item
        cells[1].text = expected
        cells[2].text = "待远端验证"
    doc.add_paragraph("该夹具不包含 API Key、用户文件或当前工作区源码。")
    doc.save(path)
    print(path)


if __name__ == "__main__":
    main()
